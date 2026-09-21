#!/usr/bin/env python3
"""docs/*.html から Word (.docx) / Excel (.xlsx) を生成する。

使い方:
    pip install python-docx openpyxl
    python3 tools/export.py            # 全ドキュメントを dist/ に出力
    python3 tools/export.py docs/sales-guideline.html

ブラウザ上のボタンからも出力できます（そちらは .doc / .xls 形式）。
より厳密な .docx / .xlsx が必要な場合にこのスクリプトを使ってください。
"""

from __future__ import annotations

import json
import re
import sys
from html.parser import HTMLParser
from pathlib import Path

from docx import Document
from docx.shared import Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
DOCS_DIR = ROOT / "docs"
DIST_DIR = ROOT / "dist"
RULES_JS = DOCS_DIR / "data" / "rules.js"
RULES_COLUMNS = ["カテゴリ", "ドキュメント", "セクション", "対象", "項目", "内容"]
RULES_KEYS = ["category", "doc", "section", "target", "item", "content"]
SKIP = {"index.html", "rules.html"}


class DocParser(HTMLParser):
    """必要な要素だけを順番に拾う簡易パーサ。"""

    BLOCKS = {"h1", "h2", "p", "li", "th", "td"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.items: list[tuple[str, str]] = []   # (kind, text)
        self.title = ""
        self._stack: list[str] = []
        self._buf: list[str] = []
        self._kind: str | None = None
        self._classes: list[str] = []
        self._in_title = False
        self._file_base = ""
        self._note_depth = 0

    def handle_starttag(self, tag, attrs):
        attr = dict(attrs)
        if tag == "title":
            self._in_title = True
        if tag == "article" and "doc" in (attr.get("class") or ""):
            self._file_base = attr.get("data-file", "")
        if tag == "div":
            classes = (attr.get("class") or "").split()
            self._note_depth += 1 if ("note" in classes or self._note_depth) else 0
        if tag == "tr":
            self.items.append(("tr-start", ""))
        if tag in self.BLOCKS:
            self._kind = tag
            self._classes = (attr.get("class") or "").split()
            self._buf = []
        self._stack.append(tag)

    def handle_endtag(self, tag):
        if tag == "title":
            self._in_title = False
        if tag == "div" and self._note_depth:
            self._note_depth -= 1
        if self._stack and self._stack[-1] == tag:
            self._stack.pop()
        if tag in self.BLOCKS and self._kind == tag:
            text = " ".join("".join(self._buf).split())
            if text:
                kind = tag
                if tag == "p" and "doc__meta" in self._classes:
                    kind = "meta"
                elif tag == "p" and "lead" in self._classes:
                    kind = "lead"
                elif tag == "p" and self._note_depth:
                    kind = "note"
                self.items.append((kind, text))
            self._kind = None
            self._buf = []

    def handle_data(self, data):
        if self._in_title:
            self.title += data
        if self._kind and "a" not in self._stack[-2:]:
            self._buf.append(data)

    @property
    def file_base(self) -> str:
        return self._file_base or self.title or "document"


def parse(path: Path) -> DocParser:
    parser = DocParser()
    parser.feed(path.read_text(encoding="utf-8"))
    return parser


def build_docx(doc_parser: DocParser, out: Path) -> None:
    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Yu Gothic"
    style.font.size = Pt(10.5)

    pending_row: list[str] = []
    table = None

    def flush_table() -> None:
        nonlocal table
        table = None

    for kind, text in doc_parser.items:
        if kind == "tr-start":
            pending_row.clear()
            continue
        if kind in ("th", "td"):
            pending_row.append(text)
            if len(pending_row) == 2:
                if table is None:
                    table = document.add_table(rows=0, cols=2)
                    table.style = "Table Grid"
                cells = table.add_row().cells
                cells[0].text, cells[1].text = pending_row
                for para in cells[0].paragraphs:
                    for run in para.runs:
                        run.bold = True
                pending_row.clear()
            continue

        flush_table()
        if kind == "h1":
            document.add_heading(text, level=1)
        elif kind == "h2":
            document.add_heading(text, level=2)
        elif kind == "li":
            document.add_paragraph(text, style="List Bullet")
        elif kind == "meta":
            para = document.add_paragraph(text)
            para.runs[0].font.size = Pt(9)
        elif kind == "note":
            para = document.add_paragraph(text)
            para.runs[0].italic = True
        else:
            document.add_paragraph(text)

    out.parent.mkdir(parents=True, exist_ok=True)
    document.save(out)


def build_xlsx(doc_parser: DocParser, out: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "内容"
    sheet.append(["区分", "項目", "内容"])

    header_fill = PatternFill("solid", fgColor="EEF4FD")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill

    section = ""
    row: list[str] = []
    for kind, text in doc_parser.items:
        if kind == "tr-start":
            row = []
        elif kind in ("th", "td"):
            row.append(text)
            if len(row) == 2:
                sheet.append([section, row[0], row[1]])
                row = []
        elif kind == "h1":
            sheet.append(["タイトル", "", text])
        elif kind == "h2":
            section = text
            sheet.append(["見出し", text, ""])
        elif kind == "li":
            sheet.append([section, "箇条書き", text])
        elif kind == "meta":
            sheet.append(["メタ情報", "", text])
        elif kind == "note":
            sheet.append([section, "注記", text])
        else:
            sheet.append([section, "本文", text])

    for column, width in zip("ABC", (26, 22, 80)):
        sheet.column_dimensions[column].width = width
    for line in sheet.iter_rows(min_row=1):
        for cell in line:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(3)}{sheet.max_row}"

    out.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(out)


def load_rules() -> list[dict]:
    """docs/data/rules.js の window.RULES を読み込む。"""
    text = RULES_JS.read_text(encoding="utf-8")
    match = re.search(r"window\.RULES\s*=\s*(\[.*\])\s*;", text, re.S)
    if not match:
        raise ValueError(f"{RULES_JS} から window.RULES を読み取れませんでした。")
    return json.loads(match.group(1))


def build_rules_docx(rules: list[dict], out: Path) -> None:
    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Yu Gothic"
    style.font.size = Pt(9)

    document.add_heading("ルール一覧", level=1)
    table = document.add_table(rows=1, cols=len(RULES_COLUMNS))
    table.style = "Table Grid"
    for cell, label in zip(table.rows[0].cells, RULES_COLUMNS):
        cell.text = label
        for para in cell.paragraphs:
            for run in para.runs:
                run.bold = True

    for rule in rules:
        cells = table.add_row().cells
        for cell, key in zip(cells, RULES_KEYS):
            cell.text = rule.get(key, "")

    out.parent.mkdir(parents=True, exist_ok=True)
    document.save(out)


def build_rules_xlsx(rules: list[dict], out: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "ルール一覧"
    sheet.append(RULES_COLUMNS)

    header_fill = PatternFill("solid", fgColor="EEF4FD")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill

    for rule in rules:
        sheet.append([rule.get(key, "") for key in RULES_KEYS])

    for column, width in zip("ABCDEF", (20, 30, 34, 26, 24, 70)):
        sheet.column_dimensions[column].width = width
    for line in sheet.iter_rows(min_row=1):
        for cell in line:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(RULES_COLUMNS))}{sheet.max_row}"

    out.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(out)


def main(argv: list[str]) -> int:
    targets = [Path(a) for a in argv[1:]] or sorted(
        p for p in DOCS_DIR.glob("*.html") if p.name not in SKIP
    )
    if not targets:
        print("出力対象の HTML が見つかりませんでした。", file=sys.stderr)
        return 1

    for path in targets:
        parsed = parse(path)
        base = parsed.file_base
        build_docx(parsed, DIST_DIR / f"{base}.docx")
        build_xlsx(parsed, DIST_DIR / f"{base}.xlsx")
        print(f"{path.name} -> dist/{base}.docx, dist/{base}.xlsx")

    if not argv[1:] and RULES_JS.exists():
        rules = load_rules()
        build_rules_docx(rules, DIST_DIR / "ルール一覧.docx")
        build_rules_xlsx(rules, DIST_DIR / "ルール一覧.xlsx")
        print(f"rules.js ({len(rules)}件) -> dist/ルール一覧.docx, dist/ルール一覧.xlsx")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
